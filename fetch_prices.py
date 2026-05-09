import yfinance as yf
import pandas as pd
from datetime import datetime
import sys
import os
from openpyxl import Workbook, load_workbook
from openpyxl.utils.dataframe import dataframe_to_rows
from config import TICKERS, PRICE_PERIOD, PRICE_START_DATE, PRICE_END_DATE, PRICE_INTERVAL, OUTPUT_FILE


def fetch_historical_prices(tickers: list[str], period: str = None, start: str = None, end: str = None, interval: str = "1d") -> dict[str, pd.DataFrame]:
    results = {}
    for symbol in tickers:
        print(f"Fetching {symbol}...")
        ticker = yf.Ticker(symbol)
        if start:
            hist = ticker.history(start=start, end=end, interval=interval)
        else:
            hist = ticker.history(period=period or "1y", interval=interval)

        if hist.empty:
            print(f"  WARNING: No data returned for {symbol}. Check the ticker symbol.")
            continue

        hist.index = hist.index.tz_localize(None)  # strip timezone for Excel compatibility
        hist = hist[["Open", "Close", "High", "Low", "Volume"]]
        hist.index.name = "Date"
        results[symbol] = hist
        print(f"  {len(hist)} rows retrieved ({hist.index[0].date()} to {hist.index[-1].date()})")

    return results


def write_to_excel(data: dict[str, pd.DataFrame], output_path: str):
    # Load or create workbook
    if os.path.exists(output_path):
        wb = load_workbook(output_path)
    else:
        wb = Workbook()
        wb.remove(wb.active)  # Remove default sheet

    # Write each ticker's data
    for symbol, df in data.items():
        sheet_name = symbol[:31]  # Excel sheet names max 31 chars

        # Remove sheet if it exists
        if sheet_name in wb.sheetnames:
            del wb[sheet_name]

        # Create new sheet and write data
        ws = wb.create_sheet(sheet_name)
        for r_idx, row in enumerate(dataframe_to_rows(df, index=True, header=True), 1):
            for c_idx, value in enumerate(row, 1):
                ws.cell(row=r_idx, column=c_idx, value=value)

        # Format columns
        ws.column_dimensions["A"].width = 18  # Date
        for col in ["B", "C", "D", "E", "F"]:
            ws.column_dimensions[col].width = 14

        # Format numeric columns
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
            for cell in row[1:5]:  # Open, Close, High, Low
                cell.number_format = "#,##0.00"
            row[5].number_format = "#,##0"  # Volume

    wb.save(output_path)
    print(f"\nSaved to: {os.path.abspath(output_path)}")


if __name__ == "__main__":
    if PRICE_PERIOD:
        data = fetch_historical_prices(TICKERS, period=PRICE_PERIOD, interval=PRICE_INTERVAL)
    else:
        data = fetch_historical_prices(TICKERS, start=PRICE_START_DATE, end=PRICE_END_DATE, interval=PRICE_INTERVAL)

    if not data:
        print("No data fetched. Exiting.")
        sys.exit(1)

    write_to_excel(data, OUTPUT_FILE)
